import tempfile
import unittest
from pathlib import Path

import openpyxl

from catalog_agent import extract


def _wb(sheets: dict[str, list[list]]) -> str:
    """Создаёт временный .xlsx с указанными листами (имя -> список строк) и возвращает путь."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    wb.save(path)
    return path


class TestNorm(unittest.TestCase):
    def test_norm_article(self):
        self.assertEqual(extract.norm_article("'2610A03364"), "2610A03364")
        self.assertEqual(extract.norm_article("  06019h2200 "), "06019H2200")
        self.assertIsNone(extract.norm_article(None))
        self.assertIsNone(extract.norm_article(""))

    def test_norm_label(self):
        self.assertEqual(extract.norm_label("Артикул/ Art.No"), "артикул")
        self.assertEqual(extract.norm_label("Артикул/\nArt.No"), "артикул")
        self.assertEqual(extract.norm_label("Стандартный артикул"), "стандартный артикул")
        self.assertIsNone(extract.norm_label(3))  # числовые ячейки — не заголовки

    def test_clean_unit(self):
        self.assertEqual(extract.clean_unit("ШТ, ШТ"), "ШТ")
        self.assertIsNone(extract.clean_unit(None))


class TestReadPricelist(unittest.TestCase):
    def test_two_row_header_with_segment_and_model(self):
        path = _wb(
            {
                "BLUE": [
                    ["Прайс лист:", None, None, None],
                    ["Артикул/ Art.No", "Сегмент/ Segment", "Модель/ Model", "Вес нетто, кг"],
                    ["Артикул", "Сегмент", "Модель", "Вес нетто, кг"],
                    ["A1", "Дрели", "GBH 220", 1.5],
                ]
            }
        )
        try:
            result = extract.read_pricelist(path)
        finally:
            Path(path).unlink()

        self.assertIn("A1", result)
        rec = result["A1"]
        self.assertEqual(rec["Наименование"], "Дрели")
        self.assertEqual(rec["модель"], "GBH 220")
        self.assertEqual(rec["ВЕС нетто"], 1.5)

    def test_service_like_sheet_without_segment_falls_back_to_naimenovanie(self):
        path = _wb(
            {
                "Service": [
                    ["Артикул/\nArt.No", "Наименование/\nModel Name", "Страна производства/ Country of origin"],
                    ["S1", "Retaining ring", "CN"],
                ]
            }
        )
        try:
            result = extract.read_pricelist(path)
        finally:
            Path(path).unlink()

        self.assertIn("S1", result)
        rec = result["S1"]
        self.assertEqual(rec["модель"], "Retaining ring")
        self.assertEqual(rec["страна происхождения"], "CN")

    def test_sheet_without_artikul_column_is_skipped(self):
        path = _wb(
            {
                "Инструкция": [
                    ["Шаг", "Действие"],
                    [1, "Открыть сайт"],
                ]
            }
        )
        try:
            stats = extract.SourceStats(path=path)
            result = extract.read_pricelist(path, stats)
        finally:
            Path(path).unlink()

        self.assertEqual(result, {})
        self.assertIn("Инструкция", stats.sheets_skipped)


class TestReadCustomsRegister(unittest.TestCase):
    def test_reads_and_normalises(self):
        path = _wb(
            {
                "Лист1": [
                    ["Артикул", "Код ТНВЭД", "Страна", "Ед.изм"],
                    ["'A1", 1234567890, "КИТАЙ", "ШТ, ШТ"],
                ]
            }
        )
        try:
            result = extract.read_customs_register(path)
        finally:
            Path(path).unlink()

        self.assertIn("A1", result)
        self.assertEqual(result["A1"]["Код ТНВЭД"], 1234567890)
        self.assertEqual(result["A1"]["ЕД.измер"], "ШТ")


class TestBuildRecords(unittest.TestCase):
    def test_catalog_only_fields_preserved_and_not_overwritten(self):
        catalog = {
            "A1": {
                "артикул": "A1",
                "Сертификат": "СТ-1234",
                "товарный знак": "Bosch",
                "Код ТНВЭД": 8467292000,
                "Наименование": "старое наименование",
            }
        }
        pricelist = {
            "A1": {
                "артикул": "A1",
                "Наименование": "новое наименование",
                "модель": "GBH 220",
                "ВЕС нетто": 1.2,
            }
        }
        records, missing = extract.build_records(catalog, pricelist, {})
        self.assertEqual(len(records), 1)
        rec = records[0]
        # Тех.данные из прайс-листа перекрывают каталог
        self.assertEqual(rec["Наименование"], "новое наименование")
        self.assertEqual(rec["модель"], "GBH 220")
        # Поля, доступные только вручную, берутся из каталога и не теряются
        self.assertEqual(rec["Сертификат"], "СТ-1234")
        self.assertEqual(rec["товарный знак"], "Bosch")
        # ТН ВЭД уже был в каталоге -> не считается отсутствующим
        self.assertNotIn("A1", missing)

    def test_missing_tnved_is_flagged(self):
        pricelist = {"A2": {"артикул": "A2", "модель": "X"}}
        records, missing = extract.build_records({}, pricelist, {})
        self.assertEqual(missing, ["A2"])

    def test_customs_fills_tnved_when_catalog_has_none(self):
        pricelist = {"A3": {"артикул": "A3"}}
        customs = {"A3": {"артикул": "A3", "Код ТНВЭД": 111, "ЕД.измер": "ШТ"}}
        records, missing = extract.build_records({}, pricelist, customs)
        self.assertEqual(records[0]["Код ТНВЭД"], 111)
        self.assertEqual(records[0]["ЕД.измер"], "ШТ")
        self.assertEqual(missing, [])

    def test_articles_filter_restricts_output(self):
        catalog = {"A1": {"артикул": "A1"}, "A2": {"артикул": "A2"}}
        records, _ = extract.build_records(catalog, {}, {}, articles=["A1"])
        self.assertEqual([r["артикул"] for r in records], ["A1"])

    def test_output_columns_match_template_order(self):
        records, _ = extract.build_records({"A1": {"артикул": "A1"}}, {}, {})
        self.assertEqual(list(records[0].keys()), extract.TEMPLATE_COLUMNS)


class TestWriteOutput(unittest.TestCase):
    def test_writes_header_and_rows_and_review_sheet(self):
        records = [{col: None for col in extract.TEMPLATE_COLUMNS}]
        records[0]["артикул"] = "A1"
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        try:
            extract.write_output(records, ["A1"], path)
            wb = openpyxl.load_workbook(path)
            ws = wb["Каталог"]
            self.assertEqual([c.value for c in ws[1]], extract.TEMPLATE_COLUMNS)
            self.assertEqual(ws["A2"].value, "A1")
            self.assertIn("Проверить вручную", wb.sheetnames)
        finally:
            Path(path).unlink()


if __name__ == "__main__":
    unittest.main()
